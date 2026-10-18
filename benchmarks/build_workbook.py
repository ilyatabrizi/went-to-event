import json, math, sys
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment

root = Path(__file__).resolve().parent
source = Path(sys.argv[1]).resolve()
data = json.loads(source.read_text())
wb = Workbook(); wb.remove(wb.active)
navy='18344D'; teal='DDEFEA'; gray='EEF2F6'
def sheet(name, headers, rows, widths):
    ws=wb.create_sheet(name); ws.append(headers)
    for row in rows: ws.append(row)
    ws.freeze_panes='C2'; ws.sheet_view.showGridLines=False
    ws.auto_filter.ref=ws.dimensions
    for c in ws[1]: c.fill=PatternFill('solid',fgColor=navy); c.font=Font(name='Calibri',color='FFFFFF',bold=True); c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=34
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=Font(name='Calibri',size=11,color='176B55' if c.data_type=='f' else '243746')
            c.alignment=Alignment(vertical='center')
            if c.row%2==0: c.fill=PatternFill('solid',fgColor=gray)
            if isinstance(c.value,(float,int)): c.number_format='#,##0.00' if isinstance(c.value,float) else '#,##0'
    for col,width in widths.items(): ws.column_dimensions[col].width=width
    return ws

assumptions=sheet('Assumptions',['Setting','Value','Meaning'],[
 ['p95 response target (ms)',1000,'Illustrative acceptance target; editable.'],
 ['Maximum error rate',.01,'HTTP, timeout and response validation failures combined.'],
 ['Capacity planning headroom',.5,'Retain 50% of observed maximum local stress throughput.'],
 ['Requests per active user / second',.2,'Illustrative demand: one request every five seconds.'],
 ['Run completed (UTC)',data['generatedAt'],'Timestamp from this execution.'],
 ['Dataset',100,'Synthetic future events; repository reads from memory.'],
 ['Browsing think time (ms)',1000,'One second between requests per virtual user.'],
 ['HTTP timeout (ms)',5000,'Socket inactivity timeout; no retries.'],
 ['Load mix','60 / 30 / 10','60% event list, 30% detail, 10% HTML or asset, by worker cycle.'],
 ['Arrival ramp (seconds)',1,'Users start evenly over the first second of each stage.'],
 ['Limits','60,000 requests/stage','In-flight requests may complete after the cap/duration.'],
 ['Stop rule','>5% errors or p95 >5s','Stop advancing after a completed stage crosses either limit.'],
],{'A':38,'B':34,'C':91})
assumptions['B3'].number_format='0.0%'; assumptions['B4'].number_format='0%'
headers=['Stage','Mode','Virtual users','Target seconds','Measured seconds','Requests','Errors','Error rate','Requests/sec','Mean ms','p50 ms','p95 ms','p99 ms','Max ms','Transferred MiB','Server CPU % of one core','Server RSS MiB','Event loop p99 ms','Meets targets']
rows=[]
for s in data['stages']:
    r=len(rows)+2
    rows.append([s['id'],s['mode'],s['users'],s['seconds'],s['elapsed'],s['requests'],s['errors'],f'=G{r}/F{r}',f'=F{r}/E{r}',s['meanMs'],s['p50Ms'],s['p95Ms'],s['p99Ms'],s['maxMs'],s['bytes']/1048576,s['cpuPercent'],s['rssMB'],s['loopP99Ms'],f'=IF(AND(H{r}<=\'Assumptions\'!$B$3,L{r}<=\'Assumptions\'!$B$2),"PASS","FAIL")'])
stages=sheet('Stages',headers,rows,{chr(65+i):w for i,w in enumerate([10,16,15,17,18,15,12,15,17,15,15,15,15,15,20,25,20,22,17])})
for r in range(2,stages.max_row+1): stages.cell(r,8).number_format='0.00%'
stages.conditional_formatting.add(f'S2:S{stages.max_row}',CellIsRule(operator='equal',formula=['"FAIL"'],fill=PatternFill('solid',fgColor='FADBD8')))
stages['P1'].comment=Comment('Server child process CPU time / stage elapsed wall time. 100% is one logical core. Load generator runs on the same machine. RSS is an end-of-stage sample, not a peak. Event loop metric has 20ms sampling resolution.','Codex')

groups=defaultdict(list)
for s in data['samples']: groups[(s['stage'],s['endpoint'])].append(s)
endpoint_rows=[]
for (stage,endpoint), samples in groups.items():
    times=sorted(s['ms'] for s in samples)
    q=lambda p: times[max(0,math.ceil(len(times)*p)-1)]
    endpoint_rows.append([stage,endpoint,len(samples),sum(bool(s['error']) for s in samples),sum(times)/len(times),q(.5),q(.95),q(.99),max(times),sum(s['bytes'] for s in samples)])
sheet('Endpoints',['Stage','Endpoint','Requests','Errors','Mean ms','p50 ms','p95 ms','p99 ms','Max ms','Bytes'],endpoint_rows,{'A':10,'B':42,**{c:17 for c in 'CDEFGHIJ'}})

last=stages.max_row
summary_rows=[
 ['Scope','Local production build; anonymous HTTP reads'],
 ['Total requests',f"=SUM('Stages'!F2:F{last})"],
 ['Total errors',f"=SUM('Stages'!G2:G{last})"],
 ['Overall error rate','=B4/B3'],
 ['Largest browsing stage',"=MAX('Stages'!C2:C6)"],
 ['Largest stress stage',"=MAX('Stages'!C7:C9)"],
 ['Observed peak stress requests/sec',"=MAX('Stages'!I7:I9)"],
 ['Illustrative local planning requests/sec',"=B8*'Assumptions'!B4"],
 ['Illustrative active users at assumed demand',"=B9/'Assumptions'!B5"],
 ['Hosting purchase decision','Production capacity remains unmeasured.'],
 ['Interpretation','Passing 1,000 VUs establishes only this short local read workload.'],
 ['Next capacity check','Repeat on deployed staging with real database, authentication and booking contention.'],
 ['Optimization candidate','Largest JS chunk: 683 kB raw / 198 kB gzip from the build. Assess route splitting and browser timing.'],
]
summary=sheet('Summary',['Metric','Result'],summary_rows,{'A':48,'B':112})
summary['B5'].number_format='0.00%'
summary['B8'].comment=Comment('Maximum of measured stress-stage requests/sec; derived from raw timing/counts in Stages. Not a proven server ceiling.','Codex')
summary['B10'].comment=Comment('Illustrative arithmetic only. Do not use as a live user limit: database, auth, network and writes were not exercised.','Codex')
for r in range(11,15): summary.cell(r,2).alignment=Alignment(wrap_text=True,vertical='center'); summary.row_dimensions[r].height=34
wb._sheets.remove(summary); wb._sheets.insert(0,summary)
for title,col,anchor in [('Browsing: p95 latency (ms)',12,'A17'),('Browsing: throughput (requests/sec)',9,'A34')]:
    chart=LineChart(); chart.title=title; chart.style=13; chart.height=8; chart.width=17
    chart.add_data(Reference(stages,min_col=col,min_row=1,max_row=6),titles_from_data=True)
    chart.set_categories(Reference(stages,min_col=3,min_row=2,max_row=6)); chart.x_axis.title='Virtual users'; chart.y_axis.numFmt='#,##0'; chart.legend=None
    summary.add_chart(chart,anchor)
notes=[
 ['Target','127.0.0.1 only. No requests sent to the live site or external services in this run.'],
 ['Build','API TypeScript and web Vite production builds both passed before measurement.'],
 ['Server','Actual compiled Fastify routes with a 100-event synthetic repository injected; external services disabled.'],
 ['Static delivery','Built HTML/JS/CSS served from memory by a minimal Node HTTP server; no CDN, compression or browser cache.'],
 ['Concurrency','Virtual users are sequential request loops, not simultaneous browser instances or logged-in accounts.'],
 ['Browsing','20-second stages with 1s think time, 10/100/250/500/1,000 users; not full browser journeys.'],
 ['Stress','15-second stages, no think time, 100/1,000/2,000 workers; stage cap 60,000 requests plus in-flight requests.'],
 ['Timing','Full body download included. Elapsed time includes stagger, final request drain and last think pauses.'],
 ['Percentiles','Nearest-rank percentiles over all requests, including failed requests. Errors also reported separately.'],
 ['Warmup','Each route fetched once before the first stage; 1.5s rest between stages.'],
 ['Fixture','100 synthetic copies of existing seed events, IDs bench-0..99, future dates; no database writes.'],
 ['Exclusions','Database latency, query scaling, auth, bookings, payments, chat, emails, writes, browser rendering and internet latency.'],
 ['Resource caveat','Server and load generator share one Mac. CPU, loop lag and throughput include this contention.'],
 ['Observed stress behavior','Throughput peaked at 100 workers, then declined at 1,000 and 2,000. Connection resets occurred in the latter stages; root cause was not isolated.'],
 ['Stress duration caveat','All stress stages reached the request cap early; actual elapsed seconds are in Stages. No sustained saturation ceiling was established.'],
 ['Capacity caveat','Short stages show observed behavior, not a sustained maximum, monthly user allowance or free-tier quota.'],
 ['Decision rule','Use the Stages PASS column against editable targets. Production test and hosting usage are required to decide when to upgrade.'],
 ['Production follow-up','Use authorized staging, realistic database size and user journeys; run a longer soak and inspect DB/function/egress usage.'],
 ['Historical results','Historical live burst results are from October 14, 2026. They are separate from this fresh local execution.'],
 ['Raw evidence',str(source.relative_to(root.parent))],
]+[[key,str(value)] for key,value in data['environment'].items()]
notes_ws=sheet('Method and limits',['Item','Detail'],notes,{'A':27,'B':130})
for r in range(2,notes_ws.max_row+1): notes_ws.cell(r,2).alignment=Alignment(wrap_text=True,vertical='center'); notes_ws.row_dimensions[r].height=32
old=json.loads((root/'results.json').read_text())
oldrows=[]
for round in old['rounds']:
    for name,result in round['results'].items(): oldrows.append([old['startedAt'],round['users'],name,result['requests'],result['errors'],result['p95Ms'],str(result.get('eventRecordCounts',{})),str(result.get('cacheStatusCounts',{}))])
hist=sheet('Historical live bursts',['Run UTC (historical)','Concurrent requests','Target','Requests','Errors','p95 ms','Event records','CDN cache counts'],oldrows,{'A':30,'B':23,'C':28,'D':15,'E':14,'F':16,'G':23,'H':36})
raw=sheet('Raw requests',['Stage','Virtual user','Endpoint','Latency ms','HTTP status','Error','Response bytes'],([s['stage'],s['user'],s['endpoint'],s['ms'],s['status'],s['error'],s['bytes']] for s in data['samples']),{'A':12,'B':16,'C':44,'D':20,'E':16,'F':50,'G':22})
out=source.parent/'app-load-benchmark.xlsx'; wb.save(out)
check=load_workbook(out,read_only=True)
assert check['Raw requests'].max_row-1==sum(s['requests'] for s in data['stages'])
assert len(check.sheetnames)==7
print(out)
