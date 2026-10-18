import json,sys
from pathlib import Path
from collections import Counter
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.chart import BarChart,Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment
source=Path(sys.argv[1]).resolve();data=json.loads(source.read_text());wb=Workbook();wb.remove(wb.active)
normal=Font(name='Calibri',size=11,color='243746');formula=Font(name='Calibri',size=11,color='176B55');header=Font(name='Calibri',bold=True,color='FFFFFF')
shade=PatternFill('solid',fgColor='EEF2F6');dark=PatternFill('solid',fgColor='18344D')
def sheet(name,headers,rows,widths):
 ws=wb.create_sheet(name);ws.append(headers)
 for r in rows:ws.append(r)
 ws.freeze_panes='C2';ws.sheet_view.showGridLines=False;ws.auto_filter.ref=ws.dimensions
 for row in ws:
  for c in row:
   c.font=header if c.row==1 else formula if c.data_type=='f' else normal
   c.alignment=Alignment(vertical='center',wrap_text=c.row==1)
   if c.row==1:c.fill=dark
   elif c.row%2==0:c.fill=shade
   if isinstance(c.value,float):c.number_format='#,##0.00'
   elif isinstance(c.value,int):c.number_format='#,##0'
 for col,width in widths.items():ws.column_dimensions[col].width=width
 ws.row_dimensions[1].height=36
 return ws
summary=sheet('Summary',['Metric','Result'],[
 ['Scope','Live Vercel delivery and direct Supabase event reads'],
 ['Total measured requests',f"=SUM('Stages'!E2:E{len(data['stages'])+1})"],
 ['Total measured errors',f"=SUM('Stages'!F2:F{len(data['stages'])+1})"],
 ['Overall error rate','=B4/B3'],
 ['Tested virtual users','100/1,000 attempted; see Coverage for skipped stages'],
 ['Database dataset',data['preflight']['databaseRecords']],
 ['Deployment identity','Browser Supabase host and API/database event IDs matched'],
 ['Maximum capacity','Not established: bounded paced load, short stages, small dataset'],
 ['Server purchase decision','Current hosting quotas and server-side resource usage still need review'],
 ['Stop condition','Safety stops recorded; see Stages and Coverage' if data.get('stoppedReason') else 'No automatic stop triggered'],
 ['Completed (UTC)',data['generatedAt']],
 ['Workbook verification','Formula values and tables checked; native Excel charts not visually verified'],
],{'A':44,'B':108})
summary['B5'].number_format='0.00%'
for r in range(2,14):summary.cell(r,2).alignment=Alignment(wrap_text=True,vertical='center');summary.row_dimensions[r].height=34
assumptions=sheet('Assumptions',['Setting','Value','Meaning'],[
 ['p95 target (ms)',2000,'Illustrative performance target; editable.'],
 ['Maximum error rate',.01,'Illustrative acceptance threshold; editable.'],
 ['Request interval / VU (s)',5,'Each virtual user makes sequential requests at this interval.'],
 ['Stage target duration (s)',20,'Responses may finish after the deadline.'],
 ['User arrival ramp (s)',5,'Users arrive evenly across this period.'],
 ['Load request cap per run',18000,'Includes stage warmup calls; excludes preflight identity checks.'],
 ['Request timeout (ms)',10000,'Full fetch including body response. No retries.'],
],{'A':35,'B':20,'C':100});assumptions['B3'].number_format='0.0%'
rows=[]
for s in data['stages']:
 r=len(rows)+2
 rows.append([s['id'],s['target'],s['users'],s['elapsed'],s['requests'],s['errors'],f'=F{r}/E{r}',f'=E{r}/D{r}',s['meanMs'],s['p50Ms'],s['p95Ms'],s['p99Ms'],s['maxMs'],s['peakInFlight'],s['bytes']/1048576,json.dumps(s['statuses']),json.dumps(s['caches']),json.dumps(s['recordCounts']),f'=IF(AND(G{r}<=\'Assumptions\'!$B$3,K{r}<=\'Assumptions\'!$B$2),"PASS","FAIL")',s['runnerCpuPercent'],s.get('p95WireMs',s['p95Ms']),s.get('p95QueueMs',0),s.get('inFlightLimit'),s.get('stoppedReason')])
stages=sheet('Stages',['Stage','Target','Virtual users','Elapsed seconds','Requests','Errors','Error rate','Requests/sec','Mean ms','p50 ms','p95 ms','p99 ms','Max ms','Peak in-flight requests','Response MiB','HTTP statuses','Vercel cache states','Event record counts','Meets targets','Runner CPU % of one core','Wire p95 ms','Queue p95 ms','In-flight limit','Stop reason'],rows,{'A':10,'B':35,**{chr(c):17 for c in range(ord('C'),ord('N')+1)},'N':22,'O':19,'P':25,'Q':31,'R':26,'S':18,'T':28,'U':18,'V':18,'W':18,'X':65})
for r in range(2,stages.max_row+1):stages.cell(r,7).number_format='0.00%'
stages.conditional_formatting.add(f'S2:S{stages.max_row}',CellIsRule(operator='equal',formula=['"FAIL"'],fill=PatternFill('solid',fgColor='FADBD8')))
stages['N1'].comment=Comment('Maximum simultaneous HTTP requests observed by this load generator. Virtual user count is not the number of SQL connections.','Codex')
stages['T1'].comment=Comment('Load generator CPU only. Vercel function and Supabase CPU/memory metrics were not collected.','Codex')
notes=[
 ['Site',data['preflight']['site']],['API',data['preflight']['api']],['Supabase project host',data['preflight']['supabaseHost']],
 ['Identity checks','Browser bundle Supabase hostname matched configuration; API and direct database event IDs matched before testing.'],
 ['Direct DB workload','HTTPS PostgREST read using the same joined event/venue/ticket tier selection and discovery filters as the repository. No API/CDN cache.'],
 ['Database role','Existing server-side service-role credential, same role as the API repository; no per-user RLS/authentication workload. Credential is not stored in this workbook.'],
 ['Vercel workload','Home HTML, /demo HTML and /events API tested separately with normal cache behavior. No cache busting.'],
 ['Cache interpretation','HIT measures cached delivery. MISS means edge miss; the application may still serve its in-process repository cache. Neither is a count of SQL queries.'],
 ['Direct read interpretation','A successful direct request executes the PostgREST read path. PostgreSQL can reuse internal buffers; this is not a cold-disk benchmark.'],
 ['User simulation','100 and 1,000 virtual request loops per target, staggered over five seconds, one request per five seconds. No actual browser sessions.'],
 ['Timing','Full response-body time measured from this Mac over the internet. Includes client/network latency. Not LCP/INP or server-only execution time.'],
 ['Rate interpretation','1,000 VUs at the chosen interval offer about 200 requests/sec after the ramp. Achieved throughput is reported; no capacity ceiling is inferred.'],
 ['Connection limits','After the first safety stop, database/API tests limit simultaneous HTTP requests to 64. Client queue time is included in overall latency; wire-only and queue p95 are separate.'],
 ['Failure interpretation','Network errors and client timeouts do not establish server overload. No server CPU/connection metrics or HTTP rejection was obtained for these errors.'],
 ['Test coverage','Skipped stages are explicitly listed in Coverage. A direct Supabase 1,000-user stage was not attempted after connection failures.'],
 ['Dataset limitation',f"Only {data['preflight']['databaseRecords']} live event(s) returned at preflight. Large catalogues, large joins and write contention were not tested."],
 ['Automatic stops','Any HTTP 429 or 5xx; after 50 completions, more than 2% errors or more than 10% requests taking over five seconds. In-flight requests drain.'],
 ['Exclusions','No bookings, writes, authentication, chat, emails, payments, or changes to deployment/database configuration.'],
 ['Metrics unavailable','No Vercel function concurrency/duration, Supabase CPU/IO/connections, billing quotas or remaining allowances measured.'],
 ['Upgrade implication','These stages measure current anonymous event-read behavior. Review actual hosting usage and repeat a longer realistic test before choosing an upgrade threshold.'],
 ['Error accounting','Non-2xx, timeouts, network errors and invalid payloads count as failures. No automatic retries. All errors remain in Raw requests and Errors.'],
 ['Percentiles','Nearest-rank latency over every completed attempt, including failures. Requests/sec includes failures.'],
 ['Load generator',str(data['runner'])],['Raw evidence',str(source.relative_to(source.parent.parent.parent))],
 ['Visual verification','A table preview was inspected. Native chart rendering unavailable in this session.'],
]
method=sheet('Method and limits',['Item','Detail'],notes,{'A':32,'B':135})
for r in range(2,method.max_row+1):method.cell(r,2).alignment=Alignment(wrap_text=True,vertical='center');method.row_dimensions[r].height=43
rawheaders=['Stage','Virtual user','Target','Started UTC','Latency ms','HTTP status','Error','Bytes','Vercel cache','Event records','Wire ms','Queue ms','Cause code']
def rawrow(s):return [s['stage'],s['user'],s['target'],s['startedAt'],s['ms'],s['status'],s['error'],s['bytes'],s['cache'],s['records'],s.get('wireMs',s['ms']),s.get('queueMs',0),s.get('causeCode')]
widths={'A':10,'B':15,'C':35,'D':29,'E':18,'F':15,'G':48,'H':17,'I':20,'J':18,'K':18,'L':18,'M':22}
sheet('Raw requests',rawheaders,(rawrow(s) for s in data['samples']),widths)
error_samples=[s for s in data['samples'] if s['error']]
sheet('Errors',rawheaders,(rawrow(s) for s in error_samples),widths)
coverage=[]
for name in ['Vercel home HTML','Vercel app HTML','Vercel events API','Supabase direct events read']:
 for users in [100,1000]:
  matched=[s for s in data['stages'] if s['target']==name and s['users']==users]
  if matched:
   s=matched[0];coverage.append([name,users,'Stopped early' if s.get('stoppedReason') else 'Completed',s['id'],s.get('stoppedReason') or 'See measured results in Stages'])
  else:coverage.append([name,users,'Not run',None,'Earlier safety stop; no result inferred for this stage'])
sheet('Coverage',['Target','Virtual users','Execution','Stage','Explanation'],coverage,{'A':36,'B':18,'C':23,'D':12,'E':85})
# Put a concise 1,000-user table directly below the decision notes.
summary.cell(16,1,'1,000-user stage');summary.cell(16,2,'p95 latency (ms)')
r=17
for i,s in enumerate(data['stages'],2):
 if s['users']==1000:summary.cell(r,1,f"='Stages'!B{i}");summary.cell(r,2,f"='Stages'!K{i}");summary.cell(r,2).number_format='#,##0.0';r+=1
if r>17:
 chart=BarChart();chart.title='Live response time at 1,000 virtual users';chart.y_axis.title='p95 latency (ms)';chart.y_axis.numFmt='#,##0';chart.style=13;chart.legend=None;chart.height=9;chart.width=22
 chart.add_data(Reference(summary,min_col=2,min_row=16,max_row=r-1),titles_from_data=True);chart.set_categories(Reference(summary,min_col=1,min_row=17,max_row=r-1));summary.add_chart(chart,'A23')
wb._sheets=[summary,stages,wb['Coverage'],assumptions,method,wb['Errors'],wb['Raw requests']]
out=source.parent/'vercel-supabase-benchmark.xlsx';wb.save(out)
check=load_workbook(out,read_only=True);assert check['Raw requests'].max_row-1==sum(s['requests'] for s in data['stages']);print(out)
